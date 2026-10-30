"""Render the validated workshop spec as editable PowerPoint objects."""

import argparse
from io import BytesIO
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

C = {
    "navy": "142D3A", "teal": "126B60", "mint": "E6F3E9", "amber": "FFF0D6",
    "neutral": "F1F4F4", "ink": "203C47", "muted": "536B74",
    "border": "D7E2E3", "white": "FFFFFF",
}


def color(value):
    return RGBColor.from_string(C.get(value, value))


def textbox(slide, value, x, y, w, h, size=20, tone="ink", bold=False,
            align=PP_ALIGN.LEFT, center=False, fill=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if center else MSO_ANCHOR.TOP
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(tone)
    return shape


def box(slide, x, y, w, h, tone, decision=False, ellipse=False):
    kind = MSO_SHAPE.DIAMOND if decision else MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = 0.1
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(tone)
    shape.line.color.rgb = color(tone if tone in {"navy", "teal"} else "border")
    shape.line.width = Pt(0.8)
    return shape


def diagram(slide, model):
    scale, ox, oy = 0.01, 0.66, 1.65
    for edge in model["edges"]:
        for index, end in enumerate(edge["points"][1:]):
            start = edge["points"][index]
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(ox + start[0] * scale), Inches(oy + start[1] * scale),
                Inches(ox + end[0] * scale), Inches(oy + end[1] * scale),
            )
            connector.line.color.rgb = color("muted")
            connector.line.width = Pt(1.2)
            if edge.get("dashed"):
                connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if index == len(edge["points"]) - 2 and edge.get("arrow", True):
                arrow = OxmlElement("a:tailEnd")
                arrow.set("type", "triangle")
                connector.line._get_or_add_ln().append(arrow)
        if edge.get("label"):
            width = min(5, max(0.4, len(edge["label"]) * 0.079 + 0.12))
            textbox(slide, edge["label"], ox + edge["labelX"] * scale - width / 2,
                    oy + edge["labelY"] * scale - 0.13, width, 0.26,
                    size=10.5, tone="muted", align=PP_ALIGN.CENTER, fill="white")
    for node in model["nodes"]:
        x, y, w, h = node["x"] * scale + ox, node["y"] * scale + oy, node["w"] * scale, node["h"] * scale
        tone = node.get("tone", "mint")
        decision = node.get("shape") == "decision"
        box(slide, x, y, w, h, tone, decision=decision)
        textbox(slide, "\n".join(node["lines"]), x + 0.06, y + 0.05, w - 0.12, h - 0.1,
                size=11 if decision else 13, tone="white" if tone in {"navy", "teal"} else "ink",
                bold=True, align=PP_ALIGN.CENTER, center=True)


def table(slide, item):
    rows = [item["headers"], *item["rows"]]
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.66), Inches(1.95), Inches(12), Inches(4.15))
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            cell = shape.table.cell(row_index, column)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = color("teal" if row_index == 0 else "neutral" if row_index % 2 else "white")
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(16)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = color("white" if row_index == 0 else "ink")


def render(spec_path, output):
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    if spec.get("synthetic") is not True or len(spec.get("slides", [])) != 8:
        raise ValueError("Expected a validated synthetic eight-slide spec.")
    models = json.loads((spec_path.parent / spec["diagramFile"]).read_text(encoding="utf-8"))
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = spec["title"]
    presentation.core_properties.author = "Banking Skills Lab"
    presentation.core_properties.subject = "SINTETIS / BUKAN PRODUKSI"
    presentation.core_properties.keywords = "synthetic, banking workshop, not production"
    for index, item in enumerate(spec["slides"]):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.12))
        strip.fill.solid()
        strip.fill.fore_color.rgb = color("teal")
        strip.line.fill.background()
        textbox(slide, "BANK NUSA DEMO  /  EXECUTIVE BRIEF", 0.66, 0.35, 11.8, 0.24, size=10.5, tone="teal", bold=True)
        textbox(slide, item["title"], 0.66, 0.88, 12, 0.64, size=27, tone="navy", bold=True)
        if item["kind"] == "title":
            box(slide, 0.66, 1.95, 12, 3.75, "mint")
            textbox(slide, item["subtitle"], 1.08, 2.45, 10.9, 1.75, size=27, tone="navy", bold=True)
            textbox(slide, "Proses bisnis  /  Arsitektur  /  Keputusan manusia", 1.08, 4.65, 10.9, 0.4, size=17, tone="teal")
        elif item["kind"] == "bullets":
            for i, bullet in enumerate(item["bullets"]):
                box(slide, 0.74, 2.02 + i * 0.86, 0.12, 0.12, "teal", ellipse=True)
                textbox(slide, bullet, 1.08, 1.85 + i * 0.86, 11.3, 0.71, center=True)
        elif item["kind"] == "metrics":
            for i, metric in enumerate(item["metrics"]):
                x = 0.66 + i * 4.08
                box(slide, x, 2.05, 3.83, 3.5, "mint" if i == 1 else "neutral")
                textbox(slide, metric["value"], x + 0.24, 2.35, 3.34, 0.9, size=43, tone="teal", bold=True)
                textbox(slide, metric["label"], x + 0.24, 3.5, 3.34, 0.75, size=19, tone="navy", bold=True)
                textbox(slide, metric["detail"], x + 0.24, 4.52, 3.34, 0.7, size=14, tone="muted")
        elif item["kind"] == "diagram":
            diagram(slide, models[item["diagramKey"]])
        elif item["kind"] == "table":
            table(slide, item)
        else:
            raise ValueError(f"Unsupported slide type: {item['kind']}")
        textbox(slide, "Sumber: " + " | ".join(item["sources"]), 0.66, 6.7, 12, 0.23, size=8.5, tone="muted")
        textbox(slide, "SINTETIS / BUKAN PRODUKSI - bukan ketentuan regulator", 0.66, 7.08, 11.3, 0.2, size=8.5, tone="muted")
        textbox(slide, f"{index + 1:02}", 12.02, 7.06, 0.64, 0.24, size=10, tone="teal", align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = (
            "SINTETIS / BUKAN PRODUKSI\n" + item["notes"] + "\nSumber: " + "; ".join(item["sources"])
        )
    buffer = BytesIO()
    presentation.save(buffer)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as target:
        target.write(buffer.getvalue())


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--spec", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    render(args.spec.resolve(), args.out.resolve())


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError, KeyError, TypeError) as error:
        print(f"PowerPoint export failed: {error}", file=sys.stderr)
        sys.exit(1)
